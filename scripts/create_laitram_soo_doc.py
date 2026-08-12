#!/usr/bin/env python3
"""Generate Laitram Dual Source Redundancy Sequence of Operations Word document."""

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# Brand / presentation colors (industrial engineering — slate + teal, not purple)
NAVY = RGBColor(0x1A, 0x2B, 0x3C)
TEAL = RGBColor(0x0D, 0x6E, 0x6E)
ACCENT = RGBColor(0xC4, 0x5C, 0x26)
GRAY = RGBColor(0x4A, 0x55, 0x68)
LIGHT_TEAL = "0D6E6E"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"


def set_cell_shading(cell, hex_color: str) -> None:
    """Set cell background color."""
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_run_font(run, name="Calibri", size=11, bold=False, color=None, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), LIGHT_TEAL)
    pBdr.append(bottom)
    pPr.append(pBdr)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = GRAY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for level, size in ((1, 18), (2, 14), (3, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.bold = True
        style.font.size = Pt(size)
        style.font.color.rgb = NAVY if level == 1 else TEAL
        style.paragraph_format.space_before = Pt(16 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(6)


def add_cover(doc: Document) -> None:
    for _ in range(2):
        doc.add_paragraph()

    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = brand.add_run("QDS SYSTEMS")
    set_run_font(r, size=14, bold=True, color=TEAL)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Dual Source Electrical Redundancy")
    set_run_font(r, size=28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Automatic Open-Transition Source Transfer\nwith Closed-Transition Retransfer\n& Delayed Generator Backup")
    set_run_font(r, size=16, color=TEAL)

    add_horizontal_line(doc.add_paragraph())

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("SEQUENCE OF OPERATIONS\nClient Presentation Document")
    set_run_font(r, size=12, bold=True, color=NAVY)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    today = datetime.now().strftime("%B %d, %Y")
    r = info.add_run(
        f"Project: Laitram — Dual Power Systems A & B\n"
        f"Utility Service: 3,000 kVA | 13.8 kV / 480 V | 3Ø Pad-Mount Transformers\n"
        f"Gear Fault Rating: 100 kA\n"
        f"Document Date: {today}\n"
        f"Status: For Client Review & Discussion"
    )
    set_run_font(r, size=11, color=GRAY)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = note.add_run(
        "\nThis document describes the proposed operating sequence, detection criteria, "
        "recommended timing, capabilities, limitations, and open questions for client confirmation."
    )
    set_run_font(r, size=10, italic=True, color=GRAY)

    doc.add_page_break()


def add_toc_slide(doc: Document) -> None:
    doc.add_heading("Contents", level=1)
    add_horizontal_line(doc.add_paragraph())

    items = [
        ("1.", "Purpose & Scope"),
        ("2.", "Existing System Description"),
        ("3.", "Proposed Operating Philosophy"),
        ("4.", "Source Failure Detection Criteria"),
        ("5.", "Recommended Timing Setpoints"),
        ("6.", "Detailed Sequence of Operations"),
        ("7.", "System Modes & Interlocks"),
        ("8.", "Capabilities & Limitations"),
        ("9.", "Engineering Considerations & Risks"),
        ("10.", "Utility Approval & Compliance"),
        ("11.", "Questions for Client Confirmation"),
        ("12.", "Next Steps"),
        ("Appendix A.", "Timing Summary Table"),
        ("Appendix B.", "Glossary"),
    ]
    for num, title in items:
        p = doc.add_paragraph()
        r = p.add_run(f"{num}  {title}")
        set_run_font(r, size=12, color=NAVY)
        p.paragraph_format.space_after = Pt(4)

    doc.add_page_break()


def add_section_header(doc: Document, number: str, title: str) -> None:
    doc.add_heading(f"{number}  {title}", level=1)
    add_horizontal_line(doc.add_paragraph())


def add_body(doc: Document, text: str, bold=False, italic=False) -> None:
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, size=11, bold=bold, italic=italic, color=GRAY)


def add_bullet(doc: Document, text: str, level=0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=11, color=GRAY)
    if level:
        p.paragraph_format.left_indent = Inches(0.5 * (level + 1))


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=11, color=GRAY)


def add_callout(doc: Document, label: str, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, LIGHT_GRAY)
    p = cell.paragraphs[0]
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10, bold=True, color=ACCENT)
    r2 = p.add_run(text)
    set_run_font(r2, size=10, color=GRAY)
    doc.add_paragraph()


def add_timing_table(doc: Document) -> None:
    headers = ["Function", "Recommended Delay", "Typical Range", "Rationale"]
    rows = [
        [
            "Source failure recognition (undervoltage / phase loss)",
            "Instantaneous sense + 1.0–3.0 s confirm",
            "0.5–5 s",
            "Ignore utility blinks / momentary dips; confirm sustained loss",
        ],
        [
            "Open-transition transfer to alternate transformer (primary path)",
            "3–5 s after confirmed failure",
            "2–10 s",
            "Allow confirm window; complete alternate-source transfer before generator start",
        ],
        [
            "Dead-bus / transfer complete verification",
            "0.5–1.0 s",
            "0.2–2 s",
            "Confirm preferred source open and alternate ready before closing tie/main",
        ],
        [
            "Generator start inhibit / delay (wait for A↔B transfer)",
            "10–15 s after source failure",
            "8–30 s",
            "Generator only if alternate-source transfer fails or both utilities unavailable",
        ],
        [
            "Generator transfer (ATS open/closed as existing)",
            "Per existing ATS settings after start",
            "Site-specific",
            "Preserve current gen protection path as secondary backup",
        ],
        [
            "Utility return stability before closed-transition retransfer",
            "5–15 minutes",
            "1–30 min",
            "Confirm restored source is stable; avoid chase on intermittent restoration",
        ],
        [
            "Closed-transition parallel / retransfer window",
            "≤100 ms typical (controller-limited)",
            "≤100–150 ms",
            "Minimize paralleling of dissimilar sources; meet utility / gear constraints",
        ],
        [
            "Post-retransfer settle / return to normal",
            "30–60 s",
            "10–120 s",
            "Confirm load accepted on restored source before clearing transfer state",
        ],
    ]

    table = doc.add_table(rows=1 + len(rows), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_TEAL)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            if ri % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, size=8, color=GRAY)


def build() -> Path:
    doc = Document()
    style_doc(doc)
    add_cover(doc)
    add_toc_slide(doc)

    # 1 Purpose
    add_section_header(doc, "1", "Purpose & Scope")
    add_body(
        doc,
        "This Sequence of Operations (SOO) defines the proposed automatic control strategy for the "
        "client’s dual secondary electrical systems (Source A and Source B). The objective is to "
        "make automatic open-transition transfer between the two utility-fed transformers the "
        "primary response to loss of either secondary source, with automatic closed-transition "
        "retransfer when the failed source returns. Transition to the single standby generator "
        "shall be delayed and used only when the alternate-transformer transfer does not restore "
        "power to the protected loads.",
    )
    add_body(doc, "In scope:", bold=True)
    add_bullet(doc, "Automatic detection of secondary source failure / phase loss on Systems A and B")
    add_bullet(doc, "Automatic open-transition transfer via the secondary tie breaker")
    add_bullet(doc, "Automatic closed-transition retransfer to the restored utility source")
    add_bullet(doc, "Coordination / delay of existing generator automatic transfer logic")
    add_bullet(doc, "Identification of capabilities, limitations, approvals, and client decisions")
    add_body(doc, "Out of scope (unless separately authorized):", bold=True)
    add_bullet(doc, "Primary (13.8 kV) switching or utility primary reconfiguration")
    add_bullet(doc, "Addition of a second generator or dual-generator paralleling")
    add_bullet(doc, "Complete redesign of existing gear beyond controls, sensing, and interlocking needed for this SOO")

    # 2 Existing
    add_section_header(doc, "2", "Existing System Description")
    add_body(doc, "As understood for this proposal:", bold=True)
    add_numbered(
        doc,
        "Utility feeds each customer with a 3,000 kVA, 13.8 kV to 480 V, three-phase pad-mounted "
        "transformer with fused disconnect (System A and System B).",
    )
    add_numbered(
        doc,
        "Each transformer secondary has electrical gear including automatic transfer to generator "
        "on power loss, with closed-transition retransfer capability.",
    )
    add_numbered(
        doc,
        "There is a single generator. Operations personnel select which transformer’s loads are "
        "protected by the generator.",
    )
    add_numbered(
        doc,
        "A secondary-side tie breaker between the two transformers allows manual transfer of loads "
        "to one transformer or the other, with manual closed-transition retransfer.",
    )
    add_numbered(doc, "Gear is rated 100 kA for fault current.")
    add_callout(
        doc,
        "Baseline change",
        "Today, A↔B load shifting via the tie is manual. The requested modification automates "
        "A↔B open-transition transfer and closed-transition retransfer, and delays generator "
        "transfer so alternate-source transfer is the primary protection path.",
    )

    # 3 Philosophy
    add_section_header(doc, "3", "Proposed Operating Philosophy")
    doc.add_heading("3.1 Priority of Response", level=2)
    add_numbered(doc, "Primary: Automatic open-transition transfer from failed source to healthy alternate transformer via secondary tie.")
    add_numbered(doc, "Secondary: If alternate-source transfer fails (or both utilities are unavailable), proceed to generator start / ATS transfer for the selected protected bus.")
    add_numbered(doc, "Restoration: When the failed utility source returns and is stable, automatically retransfer closed-transition to normal (split) configuration, subject to utility approval and interlocking.")

    doc.add_heading("3.2 Transition Types", level=2)
    add_bullet(doc, "Failure transfer (A→B or B→A): Open transition (break-before-make). Preferred source opens; after dead-bus / ready checks, tie (and/or alternate main as required by scheme) closes to feed loads from the healthy transformer.")
    add_bullet(doc, "Return to normal: Closed transition (make-before-break), briefly paralleling restored utility with the alternate path, then opening the tie / returning to normal breaker lineup — only when synch / voltage / phase-angle criteria are met and utility permission conditions are satisfied.")
    add_bullet(doc, "Generator path: Remains available but intentionally delayed so it does not race the A↔B transfer.")

    doc.add_heading("3.3 Normal Configuration (Assumed)", level=2)
    add_body(
        doc,
        "In normal operation, Transformer A feeds Bus A and Transformer B feeds Bus B; the secondary "
        "tie breaker is open. Each bus may retain its existing generator ATS path, with one generator "
        "selectable to protect either A or B loads as today.",
    )
    add_callout(
        doc,
        "Capacity note",
        "When one transformer feeds both buses through the tie, that transformer and its gear must "
        "carry the combined load. Automatic transfer must be inhibited or shed loads if combined "
        "demand exceeds the healthy transformer / feeder rating.",
    )

    doc.add_page_break()

    # 4 Detection
    add_section_header(doc, "4", "Source Failure Detection Criteria")
    add_body(
        doc,
        "Source health shall be determined from three-phase sensing on each secondary source "
        "(typically PT/VT inputs to the automatic transfer / PLC / protective relay controller). "
        "Recommended criteria for declaring Source A or Source B “failed”:",
    )

    doc.add_heading("4.1 Undervoltage (Power Loss)", level=2)
    add_bullet(doc, "Any phase voltage below approximately 80–85% of nominal (≈384–408 V on a 480 V system), sustained for the failure-confirm timer.")
    add_bullet(doc, "Optional dual thresholds: pickup at ~90% for alarm; transfer initiate at ~80–85% after confirm delay.")
    add_bullet(doc, "All three phases monitored independently.")

    doc.add_heading("4.2 Phase Loss / Single-Phasing", level=2)
    add_bullet(doc, "Any one phase below undervoltage threshold while others remain near nominal.")
    add_bullet(doc, "Phase-angle / negative-sequence or voltage-unbalance detection (recommended): voltage unbalance exceeding ~5–10% for the confirm timer.")
    add_bullet(doc, "Phase rotation / loss-of-phase logic in the controller shall treat single-phasing as a transfer-worthy failure to protect motors and sensitive loads.")

    doc.add_heading("4.3 Overvoltage / Frequency (Inhibit / Optional Transfer)", level=2)
    add_bullet(doc, "Overvoltage above ~110–115% of nominal: inhibit closed-transition retransfer; optional transfer-away if sustained and alternate source is healthy.")
    add_bullet(doc, "Frequency outside approximately 57–63 Hz (60 Hz system): treat as unhealthy source for retransfer; transfer-away only if alternate source is healthy and policy requires it.")

    doc.add_heading("4.4 What Is NOT Immediate Transfer", level=2)
    add_bullet(doc, "Utility voltage dips shorter than the confirm timer (momentary blinks).")
    add_bullet(doc, "Planned maintenance with transfer placed in Manual / Test / Bypass.")
    add_bullet(doc, "Downstream feeder faults cleared by local breakers (source voltage remains healthy).")
    add_bullet(doc, "Conditions where the alternate source is also unhealthy, overloaded, or locked out.")

    add_callout(
        doc,
        "Sensing recommendation",
        "Use true three-phase sensing on each transformer secondary ahead of the main / ATS "
        "point of transfer, plus breaker status (52a/52b), lockout (86), and overload / thermal "
        "inhibits. Do not rely on a single-phase monitor alone.",
    )

    # 5 Timing
    add_section_header(doc, "5", "Recommended Timing Setpoints")
    add_body(
        doc,
        "Final setpoints are subject to site commissioning, utility requirements, existing ATS "
        "settings, and motor/process ride-through needs. The following are recommended starting "
        "values for client discussion:",
    )
    add_timing_table(doc)
    doc.add_paragraph()
    add_callout(
        doc,
        "Key coordination rule",
        "Alternate-source open transfer must complete (or be declared failed) before generator "
        "start is released. Example: failure confirm 2 s + transfer allow 5 s + margin ⇒ generator "
        "start delay ≈ 10–15 s from initial failure.",
    )

    doc.add_page_break()

    # 6 SOO detail
    add_section_header(doc, "6", "Detailed Sequence of Operations")

    doc.add_heading("6.1 Normal Operation", level=2)
    add_numbered(doc, "Main A closed; Main B closed; Tie open (or equivalent normal lineup).")
    add_numbered(doc, "Controllers continuously monitor Source A and Source B voltage/phase health and breaker status.")
    add_numbered(doc, "Generator is in automatic standby for the operator-selected protected bus; generator start remains inhibited until the A↔B transfer window expires or fails.")
    add_numbered(doc, "Closed-transition retransfer logic is armed only when sources are healthy and utility/permission interlocks are satisfied.")

    doc.add_heading("6.2 Loss of Source A (Source B Healthy) — Open Transition to B", level=2)
    add_numbered(doc, "Controller detects Source A failure per Section 4.")
    add_numbered(doc, "Failure-confirm timer starts (recommended 1–3 s). If Source A recovers before timeout, abort transfer.")
    add_numbered(doc, "On confirmed failure, verify Source B is healthy, within capacity, not locked out, and tie path is available.")
    add_numbered(doc, "If checks fail → declare Alternate Transfer Fail; release generator path for the selected bus (Section 6.4).")
    add_numbered(doc, "If checks pass → initiate open transition: trip/open Main A (preferred source) after any programmed delay (total typically 3–5 s from confirmed failure).")
    add_numbered(doc, "Verify Main A open / Bus A dead or ready for alternate feed (dead-bus check as applicable).")
    add_numbered(doc, "Close Tie breaker to feed Bus A from Source B / Bus B.")
    add_numbered(doc, "Confirm tie closed and Bus A voltage restored from Source B.")
    add_numbered(doc, "Hold generator start for remaining delay window; if transfer successful, cancel generator start for this event (unless selected bus still requires gen for other reasons).")
    add_numbered(doc, "Annunciate: “On Alternate Source — Fed via Tie from B.” Log event.")

    doc.add_heading("6.3 Loss of Source B (Source A Healthy) — Open Transition to A", level=2)
    add_body(doc, "Mirror of Section 6.2 with A/B roles reversed: open Main B, close Tie, feed Bus B from Source A, inhibit/cancel generator if transfer succeeds.")

    doc.add_heading("6.4 Both Sources Failed, or Alternate Transfer Unsuccessful — Generator Path", level=2)
    add_numbered(doc, "If both Source A and Source B are unhealthy, or alternate-source checks fail, or open transfer does not restore voltage within the transfer-fail timer:")
    add_numbered(doc, "Release generator start after the coordinated delay (recommended 10–15 s from initial failure, or immediately on dual-source failure if no healthy alternate exists).")
    add_numbered(doc, "Existing ATS / generator closed- or open-transition logic proceeds for the operator-selected protected transformer loads.")
    add_numbered(doc, "Loads on the non-selected bus remain without generator backup (existing single-generator limitation).")
    add_numbered(doc, "When utility returns on the selected bus, existing closed-transition retransfer to utility applies, then A↔B return-to-normal logic may proceed if applicable.")

    doc.add_heading("6.5 Return of Failed Utility — Automatic Closed-Transition Retransfer", level=2)
    add_numbered(doc, "Failed source voltage returns within acceptable band on all phases; frequency and rotation OK.")
    add_numbered(doc, "Source-stable timer starts (recommended 5–15 minutes). Any dropout restarts the timer.")
    add_numbered(doc, "Verify utility closed-transition permission / permissive interlocks are true.")
    add_numbered(doc, "Verify synch conditions: voltage match, frequency match, phase-angle within controller limits.")
    add_numbered(doc, "Verify available fault duty and interlocking allow brief paralleling (see Section 9).")
    add_numbered(doc, "Close the restored main breaker (parallel via closed transition).")
    add_numbered(doc, "Within the closed-transition window (typically ≤100 ms controller-limited), open the Tie breaker to return to normal split configuration.")
    add_numbered(doc, "Confirm normal lineup (Main A closed, Main B closed, Tie open) and healthy voltages.")
    add_numbered(doc, "Clear alternate-source annunciation; restore generator inhibit logic to normal automatic standby.")

    doc.add_heading("6.6 Failure During Retransfer", level=2)
    add_bullet(doc, "If synch conditions are not met within a retry window: remain on alternate source; alarm “Retransfer Inhibited — Sync Fail”; allow manual closed- or open-transition return.")
    add_bullet(doc, "If closed transition is not permitted by utility or site policy: fall back to open-transition retransfer after an intentional dead-bus delay (client decision).")
    add_bullet(doc, "If paralleling protection or breaker fail occurs: trip to a safe open-transition state; lock out automatic closed retransfer until reset.")

    doc.add_page_break()

    # 7 Modes
    add_section_header(doc, "7", "System Modes & Interlocks")
    doc.add_heading("7.1 Operating Modes", level=2)
    add_bullet(doc, "Automatic: Full SOO as described.")
    add_bullet(doc, "Manual: Operator controls mains/tie; automatic transfer inhibited.")
    add_bullet(doc, "Test / Bypass: Per existing gear provisions; automatic actions inhibited or simulated.")
    add_bullet(doc, "Generator Select A or B: Existing selection of which bus the single generator protects.")

    doc.add_heading("7.2 Minimum Interlocks (Recommended)", level=2)
    add_bullet(doc, "Do not close Tie if both mains are closed unless in an approved closed-transition sequence.")
    add_bullet(doc, "Do not close a main onto a bus already energized from the opposite source except during approved closed-transition retransfer.")
    add_bullet(doc, "Inhibit automatic tie close on alternate-source overload / overcurrent / thermal alarm.")
    add_bullet(doc, "Inhibit automatic transfer if gear lockout (86), maintenance switch, or doors/racking interlocks require it.")
    add_bullet(doc, "Inhibit closed-transition retransfer without utility permissive (where required).")
    add_bullet(doc, "Breaker failure / incomplete sequence → abort to safe state and alarm.")
    add_bullet(doc, "Prevent simultaneous competing commands between A↔B controller and generator ATS (priority: A↔B first, then gen).")

    # 8 Capabilities
    add_section_header(doc, "8", "Capabilities & Limitations")
    doc.add_heading("8.1 Capabilities After Modification", level=2)
    add_bullet(doc, "Automatic restoration of loads on a failed bus from the alternate transformer without waiting for operator manual tie operation.")
    add_bullet(doc, "Reduced reliance on the single generator for single-utility outages when the alternate transformer is healthy and has capacity.")
    add_bullet(doc, "Automatic return to normal split configuration via closed-transition retransfer when the failed utility returns and is stable.")
    add_bullet(doc, "Preserved generator backup as a delayed secondary path for the selected bus.")
    add_bullet(doc, "Clearer event annunciation and logged sequence for operations and post-event review.")

    doc.add_heading("8.2 Limitations the Client Must Understand", level=2)
    add_bullet(doc, "Open-transition A↔B transfer causes a brief outage on the affected bus (typically cycles to a few seconds depending on breaker speed and timers) — not a seamless UPS-style transfer.")
    add_bullet(doc, "One transformer must carry combined A+B load while on tie; if load exceeds rating, transfer must be blocked or loads shed — otherwise transformer / feeder overload or upstream fuse operation may occur.")
    add_bullet(doc, "Single generator still protects only the selected bus; the non-selected bus has no generator coverage during dual-utility loss.")
    add_bullet(doc, "Closed-transition retransfer briefly parallels two utility sources through customer gear; this requires utility approval and verification that available fault current during paralleling remains within the 100 kA gear rating and breaker interrupting capability.")
    add_bullet(doc, "If both utilities are lost (common-mode utility event), A↔B transfer cannot help; only generator (selected bus) can sustain load.")
    add_bullet(doc, "Downstream faults, incorrect settings, or failed breakers can prevent successful automatic transfer.")
    add_bullet(doc, "Motors and process equipment may still drop out during the open-transition dead time unless UPS / ride-through is provided separately.")
    add_bullet(doc, "Automation does not remove the need for periodic testing, maintenance, and operator training.")

    doc.add_page_break()

    # 9 Engineering considerations
    add_section_header(doc, "9", "Engineering Considerations & Risks")
    doc.add_heading("9.1 Fault Interrupting Rating at Time of Retransfer", level=2)
    add_body(
        doc,
        "Gear is stated as 100 kA rated. During closed-transition retransfer, Sources A and B may be "
        "briefly paralleled. Available fault current on the paralleled bus can approach the sum of "
        "contributions from both transformers (and any generator contribution if still connected), "
        "which may exceed the fault current present in normal split operation.",
    )
    add_bullet(doc, "A short-circuit study update is required for the paralleled condition.")
    add_bullet(doc, "Confirm breaker interrupting ratings, bus bracing, and protective device settings for the parallel window.")
    add_bullet(doc, "If paralleled available fault current would exceed 100 kA or device ratings, closed-transition automatic retransfer must not be enabled; use open-transition retransfer or utility-approved mitigation.")
    add_bullet(doc, "Protective relays must be coordinated so a fault during the brief parallel does not cascade incorrectly.")

    doc.add_heading("9.2 Transformer & Feeder Capacity", level=2)
    add_bullet(doc, "Confirm each 3,000 kVA transformer and secondary gear can support worst-case combined load, including motor inrush after outage.")
    add_bullet(doc, "Define load-shed or transfer-inhibit thresholds (kW/kVA or current).")
    add_bullet(doc, "Review primary fused disconnect ratings and utility primary constraints when one secondary carries both buses.")

    doc.add_heading("9.3 Existing Generator ATS Coordination", level=2)
    add_bullet(doc, "Existing ATS undervoltage / start timers must be lengthened or supervised so they do not start the generator before A↔B transfer completes.")
    add_bullet(doc, "Avoid conflicting close commands between tie control and ATS.")
    add_bullet(doc, "Define behavior if generator is already running / paralleled when utility events occur.")

    doc.add_heading("9.4 Grounding, Neutrals, and Circulating Currents", level=2)
    add_bullet(doc, "Review grounding scheme when briefly paralleling two utility secondaries.")
    add_bullet(doc, "Closed transition must limit circulating current via sync check; excessive phase-angle error shall block close.")

    doc.add_heading("9.5 Human Factors & Annunciation", level=2)
    add_bullet(doc, "Provide clear HMI/annunciator states: Normal / On Tie from A / On Tie from B / Retransfer Pending / Gen Active / Inhibited.")
    add_bullet(doc, "Provide local Manual override and emergency open of tie.")
    add_bullet(doc, "Train operators on single-generator selection implications after automation.")

    # 10 Utility
    add_section_header(doc, "10", "Utility Approval & Compliance")
    add_body(
        doc,
        "Closed-transition retransfer that momentarily parallels customer-owned sources tied to "
        "utility feeders typically requires utility review and written approval. Items commonly required:",
    )
    add_bullet(doc, "Interconnection / paralleling request describing closed-transition window and sync-check method")
    add_bullet(doc, "One-line diagrams showing mains, tie, ATS, and generator")
    add_bullet(doc, "Protective relay settings and interconnect protection (where applicable)")
    add_bullet(doc, "Demonstration that parallel duration is minimized and supervised")
    add_bullet(doc, "Confirmation that customer gear ratings are adequate for paralleled fault duty")
    add_bullet(doc, "Any utility requirement for permissive signal, reverse power, or anti-islanding provisions")
    add_callout(
        doc,
        "Client action",
        "Utility approval should be obtained before enabling automatic closed-transition "
        "retransfer. Until approved, the system can be commissioned for automatic open-transition "
        "A↔B transfer with manual or open-transition return-to-normal.",
    )

    doc.add_page_break()

    # 11 Questions
    add_section_header(doc, "11", "Questions for Client Confirmation")
    add_body(
        doc,
        "Please review and respond to the following so the final design and settings can be locked:",
    )

    questions = [
        "What is the normal peak and emergency load on Bus A and Bus B (kW/kVA), and the expected combined load if one transformer feeds both?",
        "Can either 3,000 kVA transformer and its secondary gear safely carry the combined load continuously? If not, which loads must auto-shed before tie close?",
        "Is the preferred failure response confirmed as: (1) auto open-transfer via tie first, (2) generator only if that fails / both utilities lost?",
        "What maximum dead-bus outage duration is acceptable for critical processes during open-transition A↔B transfer?",
        "Confirm existing generator ATS manufacturer, model, current undervoltage and start-delay settings, and closed-transition retransfer settings.",
        "Which bus is typically selected for generator protection, and how often is that selection changed?",
        "Is automatic closed-transition retransfer required, or is automatic open-transition return acceptable if utility approval or fault-duty limits block closed transition?",
        "Has the serving utility previously approved closed-transition paralleling at this site? Who is the utility contact for interconnection review?",
        "Are as-built one-lines, relay settings, breaker control schematics, and short-circuit / coordination studies available for update?",
        "What is the available utility fault contribution on each secondary (or primary) for use in a paralleled-condition short-circuit evaluation against the 100 kA rating?",
        "Are there ground-fault, zone-selective interlocking, or arc-flash constraints that affect main/tie automatic operation?",
        "Should automatic retransfer wait 5, 10, or 15 minutes (or another value) after utility return?",
        "Should operators have a “Retransfer Now” pushbutton and a “Hold on Alternate” inhibit?",
        "What annunciation / SCADA / Building Management points are required for remote monitoring?",
        "Are any life-safety, NEC Article 700/701/708, or healthcare/industrial selective-coordination requirements applicable to this gear?",
        "What commissioning window and outage allowances are available for breaker control upgrades, CT/PT verification, and end-to-end testing?",
        "Who will own final setpoint approval (client engineering, facilities, and/or insurer / AHJ)?",
        "Should the system automatically return to split-bus normal after generator retransfer to utility, or require operator confirmation?",
    ]
    for i, q in enumerate(questions, 1):
        p = doc.add_paragraph()
        r = p.add_run(f"{i}. {q}")
        set_run_font(r, size=11, color=GRAY)
        p.paragraph_format.space_after = Pt(6)

    # 12 Next steps
    add_section_header(doc, "12", "Next Steps")
    add_numbered(doc, "Client review of this Sequence of Operations and responses to Section 11 questions.")
    add_numbered(doc, "Gather one-lines, settings, load data, and existing short-circuit study.")
    add_numbered(doc, "Perform / update short-circuit evaluation for paralleled retransfer vs. 100 kA gear rating.")
    add_numbered(doc, "Engage utility for closed-transition approval path.")
    add_numbered(doc, "Develop detailed control narrative, I/O list, and interlocking logic for PLC/ATS/relay implementation.")
    add_numbered(doc, "Finalize timing setpoints; coordinate generator ATS delays.")
    add_numbered(doc, "Install, commission, and demonstrate: simulated A loss, B loss, dual loss, successful retransfer, and inhibited retransfer.")
    add_numbered(doc, "Deliver as-left settings, test records, and operator training.")

    doc.add_page_break()

    # Appendix A
    add_section_header(doc, "Appendix A", "Timing Summary (Proposed Defaults)")
    rows_a = [
        ("Failure confirm (UV / phase loss)", "2 seconds"),
        ("Open-transition transfer initiate", "3–5 seconds after confirm"),
        ("Alternate transfer fail timer", "8–10 seconds from failure"),
        ("Generator start delay", "12 seconds from failure (adjust to exceed A↔B success path)"),
        ("Utility return stable before retransfer", "10 minutes"),
        ("Closed-transition parallel window", "≤100 ms (controller limit)"),
        ("Post-retransfer normal declare", "60 seconds"),
    ]
    table = doc.add_table(rows=1 + len(rows_a), cols=2)
    table.style = "Table Grid"
    for i, h in enumerate(["Timer / Function", "Proposed Default"]):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_TEAL)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        set_run_font(r, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for ri, (a, b) in enumerate(rows_a):
        for ci, val in enumerate((a, b)):
            cell = table.rows[ri + 1].cells[ci]
            if ri % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, size=10, color=GRAY)

    doc.add_paragraph()
    add_body(
        doc,
        "All defaults are provisional pending client answers, utility requirements, and commissioning.",
        italic=True,
    )

    # Appendix B
    add_section_header(doc, "Appendix B", "Glossary")
    glossary = [
        ("Open transition", "Break-before-make transfer; load experiences a brief dead-bus interval."),
        ("Closed transition", "Make-before-break transfer; sources are briefly paralleled, then the alternate path is opened."),
        ("Tie breaker", "Secondary breaker connecting Bus A and Bus B."),
        ("ATS", "Automatic Transfer Switch (existing generator path)."),
        ("Retransfer", "Return of load from alternate (or generator) source back to the normal utility source."),
        ("Sync check", "Verification that voltage, frequency, and phase angle are within limits before paralleling."),
        ("Confirm timer", "Time a failure condition must persist before automatic action."),
        ("100 kA rated", "Gear withstand / interrupting capability stated by client for fault current evaluation."),
    ]
    for term, definition in glossary:
        p = doc.add_paragraph()
        r1 = p.add_run(f"{term}: ")
        set_run_font(r1, size=11, bold=True, color=NAVY)
        r2 = p.add_run(definition)
        set_run_font(r2, size=11, color=GRAY)

    # Closing
    doc.add_paragraph()
    add_horizontal_line(doc.add_paragraph())
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = end.add_run(
        "\nEnd of Sequence of Operations — For Client Review\n"
        "QDS Systems | Dual Source Redundancy | Laitram\n"
        "Not for construction until issued with final drawings, settings, and approvals."
    )
    set_run_font(r, size=10, italic=True, color=GRAY)

    out_dir = Path("/workspace/docs")
    out_dir.mkdir(parents=True, exist_ok=True)
    artifacts = Path("/workspace/artifacts")
    artifacts.mkdir(parents=True, exist_ok=True)

    out_path = out_dir / "Laitram_Dual_Source_Redundancy_Sequence_of_Operations.docx"
    art_path = artifacts / "Laitram_Dual_Source_Redundancy_Sequence_of_Operations.docx"
    doc.save(out_path)
    doc.save(art_path)
    return out_path


if __name__ == "__main__":
    path = build()
    print(f"Wrote: {path}")
